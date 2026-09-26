import os
import random
import shutil
import subprocess
import tempfile
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import openpyxl
from docx import Document
from docx.image.image import Image as ImageParser
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu, Inches
from docx.text.paragraph import Paragraph

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(PROJECT_DIR, "data.xlsx")
TEMPLATE_FILE = os.path.join(PROJECT_DIR, "Adventure_Template.docx")
IMAGES_DIR = os.path.join(PROJECT_DIR, "images")
OPENING_PLACEHOLDER = "[INSERT_OPENING]"
HEADING_PLACEHOLDER = "[Insert_Heading]"
IMAGE_PLACEHOLDER = "[Image_Placeholder]"
PLAYER_NAME_PLACEHOLDER = "[Player_Name]"
HEADING_TEXT = "[Player_Name]'s Journey Begins"
MAX_IMAGE_HEIGHT_INCHES = 4.5

CHARACTER_TYPES = ["Catcher"]


def get_starting_location():
    wb = openpyxl.load_workbook(DATA_FILE, read_only=True)
    ws = wb["Locations"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        location_id, location_name = row[0], row[1]
        if location_id == 1:
            return location_name
    raise ValueError("No location found with Location ID 1")


def get_location_id(location_name):
    wb = openpyxl.load_workbook(DATA_FILE, read_only=True)
    ws = wb["Locations"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] == location_name:
            return int(row[0])
    raise ValueError(f"No location named '{location_name}' on the Locations tab")


def get_location_image_path(location_id):
    """images/location_<ID>.png for the starting location's Location ID."""
    path = os.path.join(IMAGES_DIR, f"location_{location_id}.png")
    if not os.path.exists(path):
        raise FileNotFoundError(f"Missing location image: {path}")
    return path


def get_random_opening(player_type):
    wb = openpyxl.load_workbook(DATA_FILE, read_only=True)
    ws = wb["Openings"]
    matches = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        opening_text, opening_player_type = row[1], row[3]
        if opening_player_type == player_type:
            matches.append(opening_text)
    if not matches:
        raise ValueError(f"No openings found for Player_Type '{player_type}'")
    return random.choice(matches)


class ConverterMissing(Exception):
    """Raised when no docx-to-PDF converter is available on this machine."""


def _all_paragraphs(doc):
    """Every paragraph in the document body (including inside tables) plus every
    header and footer, so nothing in the template is skipped."""
    roots = [doc.element.body]
    for section in doc.sections:
        for part in (section.header, section.footer):
            roots.append(part._element)
    for root in roots:
        for p in root.iter(qn("w:p")):
            yield Paragraph(p, doc)


def _replace_in_paragraph(paragraph, old, new):
    """Replace `old` with `new` inside one paragraph while keeping the formatting
    of the run it lives in. Returns True if a replacement was made."""
    runs = paragraph.runs
    for run in runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True

    # placeholder split across several runs: merge them into the first one
    full = "".join(r.text for r in runs)
    idx = full.find(old)
    if idx == -1:
        return False
    pos = 0
    first = None
    for run in runs:
        run_start, run_end = pos, pos + len(run.text)
        pos = run_end
        if run_end <= idx or run_start >= idx + len(old):
            continue
        if first is None:
            first = run
            offset = idx - run_start
            run.text = run.text[:offset] + new + run.text[offset + (min(run_end, idx + len(old)) - idx):]
        else:
            cut_from = 0
            cut_to = min(run_end, idx + len(old)) - run_start
            run.text = run.text[cut_to:]
    return True


def _fill_opening(paragraph, opening_text):
    """Put the opening into the [INSERT_OPENING] spot. If the placeholder sits right
    after a one-letter drop cap (like the template's big blue "F"), the drop cap
    becomes the opening's first letter and the placeholder gets the rest, so the
    text doesn't read "FYou wake up...". """
    for i, run in enumerate(paragraph.runs):
        if OPENING_PLACEHOLDER not in run.text:
            continue
        prev = paragraph.runs[i - 1] if i > 0 else None
        is_drop_cap = (
            prev is not None
            and run.text == OPENING_PLACEHOLDER
            and len(prev.text) == 1
            and prev.font.size is not None
            and prev.font.size.pt >= 14
        )
        if is_drop_cap and opening_text:
            prev.text = opening_text[0]
            run.text = opening_text[1:]
        else:
            run.text = run.text.replace(OPENING_PLACEHOLDER, opening_text)
        return True
    return _replace_in_paragraph(paragraph, OPENING_PLACEHOLDER, opening_text)


def _fill_image(paragraph, image_path, max_width, max_height):
    """Swap the [Image_Placeholder] box for a picture. The placeholder paragraph is
    drawn as a dashed, shaded box, so that styling (and its indent) is removed to
    let the picture sit cleanly on the margins. The picture is always centered,
    whatever alignment the template paragraph had and whatever the image's shape."""
    p = paragraph._p
    for run_el in p.findall(qn("w:r")):
        p.remove(run_el)
    pPr = p.get_or_add_pPr()
    for tag in ("w:pBdr", "w:shd", "w:ind"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    image = ImageParser.from_file(image_path)
    width = max_width
    height = int(width * image.px_height / image.px_width)
    if height > max_height:
        height = max_height
        width = int(height * image.px_width / image.px_height)
    picture = paragraph.add_run().add_picture(image_path, width=Emu(width), height=Emu(height))
    # Without explicit zero wrap distances LibreOffice pads the picture and shifts it
    # ~0.13" right, past the margin.
    for attr in ("distT", "distB", "distL", "distR"):
        picture._inline.set(attr, "0")


def _fill_caption(image_paragraph, text):
    """Write `text` into the caption paragraph directly under the picture, keeping
    that paragraph's own formatting (the template's small right-aligned caption)."""
    following = image_paragraph._p.getnext()
    if following is None or following.tag != qn("w:p"):
        return
    caption = Paragraph(following, image_paragraph._parent)
    runs = caption.runs
    if not runs:
        caption.add_run(text)
        return
    runs[0].text = text
    for extra in runs[1:]:
        extra.text = ""


def fill_template(opening_text, player_name, location_name, image_path):
    """Open Adventure_Template.docx and fill in the placeholders in place. Every
    page, table, image, header/footer and font in the template is preserved."""
    doc = Document(TEMPLATE_FILE)
    for paragraph in list(_all_paragraphs(doc)):
        if OPENING_PLACEHOLDER in "".join(r.text for r in paragraph.runs):
            _fill_opening(paragraph, opening_text)
    for paragraph in list(_all_paragraphs(doc)):
        if HEADING_PLACEHOLDER in "".join(r.text for r in paragraph.runs):
            _replace_in_paragraph(paragraph, HEADING_PLACEHOLDER, HEADING_TEXT)

    section = doc.sections[0]
    max_width = int(section.page_width - section.left_margin - section.right_margin)
    max_height = int(Inches(MAX_IMAGE_HEIGHT_INCHES))
    for paragraph in list(_all_paragraphs(doc)):
        if IMAGE_PLACEHOLDER in "".join(r.text for r in paragraph.runs):
            _fill_image(paragraph, image_path, max_width, max_height)
            # captions in the template are all caps, so the town name matches them
            _fill_caption(paragraph, location_name.upper())

    # last, so [Player_Name] is filled in everywhere, including inside the heading
    # and the opening text
    for paragraph in list(_all_paragraphs(doc)):
        while PLAYER_NAME_PLACEHOLDER in "".join(r.text for r in paragraph.runs):
            if not _replace_in_paragraph(paragraph, PLAYER_NAME_PLACEHOLDER, player_name):
                break
    return doc


def find_libreoffice():
    found = shutil.which("soffice") or shutil.which("soffice.exe")
    if found:
        return found
    for base in (os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)")):
        if base:
            candidate = os.path.join(base, "LibreOffice", "program", "soffice.exe")
            if os.path.exists(candidate):
                return candidate
    return None


def create_adventure(opening_text, player_name, location_name, image_path, output_path):
    """Write the finished adventure. `.docx` output is the filled template itself;
    `.pdf` output is that same document converted by LibreOffice, so the layout
    and fonts match the template exactly."""
    doc = fill_template(opening_text, player_name, location_name, image_path)

    if output_path.lower().endswith(".docx"):
        doc.save(output_path)
        return

    soffice = find_libreoffice()
    if soffice is None:
        raise ConverterMissing(
            "Saving as PDF needs LibreOffice (free) to keep the template's formatting."
        )

    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_docx = os.path.join(tmp_dir, "adventure.docx")
        doc.save(tmp_docx)
        # a persistent profile: LibreOffice takes ~40s to create one, so reuse it
        profile = os.path.join(tempfile.gettempdir(), "morimon_libreoffice_profile")
        subprocess.run(
            [
                soffice,
                f"-env:UserInstallation=file:///{profile.replace(os.sep, '/')}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                tmp_dir,
                tmp_docx,
            ],
            check=True,
            timeout=180,
            capture_output=True,
        )
        tmp_pdf = os.path.join(tmp_dir, "adventure.pdf")
        if not os.path.exists(tmp_pdf):
            raise RuntimeError("LibreOffice did not produce a PDF.")
        shutil.move(tmp_pdf, output_path)


class AdventureCreatorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Morimon Story Maker")
        self.root.geometry("400x200")

        self.menu_frame = tk.Frame(self.root)
        self.form_frame = tk.Frame(self.root)

        self.show_menu()

    def show_menu(self):
        self.form_frame.pack_forget()
        for widget in self.menu_frame.winfo_children():
            widget.destroy()

        create_button = tk.Button(
            self.menu_frame,
            text="Create New Adventure",
            command=self.show_new_adventure_form,
        )
        create_button.pack(pady=20)

        self.menu_frame.pack(fill="both", expand=True)

    def show_new_adventure_form(self):
        self.menu_frame.pack_forget()
        for widget in self.form_frame.winfo_children():
            widget.destroy()

        starting_location = get_starting_location()

        tk.Label(self.form_frame, text="Character Type").grid(
            row=0, column=0, sticky="w", padx=10, pady=10
        )
        self.character_type_var = tk.StringVar(value=CHARACTER_TYPES[0])
        character_type_dropdown = ttk.Combobox(
            self.form_frame,
            textvariable=self.character_type_var,
            values=CHARACTER_TYPES,
            state="readonly",
        )
        character_type_dropdown.grid(row=0, column=1, padx=10, pady=10)

        tk.Label(self.form_frame, text="Starting Location").grid(
            row=1, column=0, sticky="w", padx=10, pady=10
        )
        self.starting_location_var = tk.StringVar(value=starting_location)
        starting_location_dropdown = ttk.Combobox(
            self.form_frame,
            textvariable=self.starting_location_var,
            values=[starting_location],
            state="readonly",
        )
        starting_location_dropdown.grid(row=1, column=1, padx=10, pady=10)

        tk.Label(self.form_frame, text="Player Name").grid(
            row=2, column=0, sticky="w", padx=10, pady=10
        )
        self.player_name_var = tk.StringVar()
        player_name_entry = tk.Entry(self.form_frame, textvariable=self.player_name_var)
        player_name_entry.grid(row=2, column=1, padx=10, pady=10)

        button_frame = tk.Frame(self.form_frame)
        button_frame.grid(row=3, column=0, columnspan=2, pady=10)

        tk.Button(button_frame, text="Create Adventure", command=self.submit_form).pack(
            side="left", padx=5
        )
        tk.Button(button_frame, text="Back", command=self.show_menu).pack(
            side="left", padx=5
        )

        self.form_frame.pack(fill="both", expand=True)

    def submit_form(self):
        Player_Name = self.player_name_var.get().strip()

        if not Player_Name:
            messagebox.showerror("Missing Player Name", "Please enter a player name.")
            return

        Player_Type = self.character_type_var.get()
        starting_location = self.starting_location_var.get()

        try:
            opening_text = get_random_opening(Player_Type)
        except ValueError as exc:
            messagebox.showerror("No Opening Found", str(exc))
            return

        try:
            image_path = get_location_image_path(get_location_id(starting_location))
        except (ValueError, FileNotFoundError) as exc:
            messagebox.showerror("Location Image Not Found", str(exc))
            return

        output_path = filedialog.asksaveasfilename(
            title="Save Adventure As",
            defaultextension=".pdf",
            initialfile=f"{Player_Name}_Adventure.pdf",
            filetypes=[("PDF files", "*.pdf"), ("Word documents", "*.docx")],
        )
        if not output_path:
            return

        try:
            create_adventure(opening_text, Player_Name, starting_location, image_path, output_path)
        except ConverterMissing as exc:
            save_docx = messagebox.askyesno(
                "PDF converter not found",
                f"{exc}\n\nSave a Word document (.docx) instead? It has the exact same "
                "pages and formatting, and you can turn it into a PDF from Word "
                "(File > Save As > PDF).",
            )
            if not save_docx:
                return
            output_path = os.path.splitext(output_path)[0] + ".docx"
            try:
                create_adventure(opening_text, Player_Name, starting_location, image_path, output_path)
            except Exception as inner_exc:
                messagebox.showerror("Failed to Create Adventure", str(inner_exc))
                return
        except Exception as exc:
            messagebox.showerror("Failed to Create Adventure", str(exc))
            return

        messagebox.showinfo(
            "Adventure Created",
            f"Player Name: {Player_Name}\n"
            f"Character Type: {Player_Type}\n"
            f"Starting Location: {starting_location}\n\n"
            f"Saved to: {output_path}",
        )


def main():
    root = tk.Tk()
    AdventureCreatorApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
