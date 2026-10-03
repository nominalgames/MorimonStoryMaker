"""Chooses and reads the player's "Main" questline, from the `Quests` tab.

A questline is a group of rows on `Quests` that share a Questline_ID. Each row is
one step: the Location ID and Sub_Location_ID it plays out at, and the Text to
show there (matched against location_reveal.SubLocation.sub_location_id).

`choose_main_quest(player_type)` is meant to run once, right after the player
enters their name and clicks Create Adventure: it picks one "Main" questline at
random from the ones written for that Player_Type, and returns all of its steps.
pdf_builder.draw_location_reveal then prints a step's Text right after that
sub-location's own Sub-Location-Desc, if this playthrough's quest has one there.
"""
import os
import random
from dataclasses import dataclass, field

import openpyxl

DATA_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data.xlsx")

QUESTS_SHEET = "Quests"
MAIN_MODIFIER = "Main"

QUESTLINE_ID_HEADER = "Questline_ID"
QUESTLINE_NAME_HEADER = "Questline_Name"
SEQUENCE_NAME_HEADER = "Sequence_Name"
LOCATION_HEADER = "Location(s)"
SUB_LOCATION_HEADER = "Sub_Location"
MODIFIERS_HEADER = "Modifiers"
PLAYER_TYPE_HEADER = "Player_Type"
TEXT_HEADER = "Text"


def _clean(value):
    return "" if value is None else str(value).strip()


def _to_int(value):
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


@dataclass
class QuestStep:
    questline_id: int
    sequence_name: str
    location_id: int
    sub_location_id: int
    text: str


@dataclass
class MainQuest:
    questline_id: int
    name: str
    steps: list = field(default_factory=list)  # list[QuestStep]

    def text_for(self, location_id, sub_location_id):
        """The Text for this quest's step at this location/sub-location, or ""
        if it doesn't have one there."""
        if location_id is None or sub_location_id is None:
            return ""
        for step in self.steps:
            if step.location_id == location_id and step.sub_location_id == sub_location_id:
                return step.text
        return ""


def choose_main_quest(player_type):
    """Pick one random "Main" questline written for `player_type` and return it
    with all of its steps. Returns None if there's no matching questline (e.g.
    the `Quests` tab is missing, empty, or has nothing for this Player_Type),
    so the adventure can be created without one."""
    wb = openpyxl.load_workbook(DATA_FILE, read_only=True, data_only=True)
    if QUESTS_SHEET not in wb.sheetnames:
        return None

    rows = list(wb[QUESTS_SHEET].iter_rows(values_only=True))
    if not rows:
        return None
    headers = [_clean(h) for h in rows[0]]
    if QUESTLINE_ID_HEADER not in headers:
        return None

    idx = {
        name: headers.index(name)
        for name in (
            QUESTLINE_ID_HEADER, QUESTLINE_NAME_HEADER, SEQUENCE_NAME_HEADER, LOCATION_HEADER,
            SUB_LOCATION_HEADER, MODIFIERS_HEADER, PLAYER_TYPE_HEADER, TEXT_HEADER,
        )
        if name in headers
    }

    def get(row, header):
        i = idx.get(header)
        return row[i] if i is not None and i < len(row) else None

    main_rows = [
        row for row in rows[1:]
        if _to_int(get(row, QUESTLINE_ID_HEADER)) is not None
        and _clean(get(row, MODIFIERS_HEADER)) == MAIN_MODIFIER
        and _clean(get(row, PLAYER_TYPE_HEADER)) == player_type
    ]
    if not main_rows:
        return None

    questline_ids = sorted({_to_int(get(row, QUESTLINE_ID_HEADER)) for row in main_rows})
    chosen_id = random.choice(questline_ids)

    # every row for the chosen questline, not just the ones that matched the
    # Modifiers/Player_Type filter above (a later step's row might leave those blank)
    steps = []
    name = ""
    for row in rows[1:]:
        if _to_int(get(row, QUESTLINE_ID_HEADER)) != chosen_id:
            continue
        name = name or _clean(get(row, QUESTLINE_NAME_HEADER))
        loc_id = _to_int(get(row, LOCATION_HEADER))
        sub_id = _to_int(get(row, SUB_LOCATION_HEADER))
        text = _clean(get(row, TEXT_HEADER))
        if loc_id is None or sub_id is None or not text:
            continue
        steps.append(QuestStep(
            questline_id=chosen_id,
            sequence_name=_clean(get(row, SEQUENCE_NAME_HEADER)),
            location_id=loc_id,
            sub_location_id=sub_id,
            text=text,
        ))

    return MainQuest(questline_id=chosen_id, name=name, steps=steps)
